"""Conservative, atomic XLSX ingestion into the shared reconciliation model."""

from collections.abc import Iterator, Sequence
from contextlib import contextmanager
from datetime import date, datetime, time, timedelta
from decimal import Decimal
from io import BytesIO
from math import isfinite
from xml.etree.ElementTree import ParseError, fromstring
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.cell.read_only import EmptyCell, ReadOnlyCell
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet._read_only import ReadOnlyWorksheet

from tallydiff.amounts import parse_amount
from tallydiff.ingest import IngestionError
from tallydiff.models import Source, SourceRecord

type Cell = ReadOnlyCell | EmptyCell


@contextmanager
def _open(data: bytes, source: Source, worksheet: str | None = None) -> Iterator[Workbook]:
    if not isinstance(source, Source):
        raise TypeError("source must be a Source enum member")
    if not isinstance(data, bytes):
        raise TypeError("data must be XLSX bytes")
    # Own the stream and close the lazy workbook even when validation fails.
    with BytesIO(data) as stream:
        try:
            with ZipFile(stream) as archive:
                manifest = fromstring(archive.read("[Content_Types].xml"))
                if not any(
                    entry.get("ContentType")
                    == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"
                    for entry in manifest
                ):
                    raise IngestionError(
                        source, "only standard .xlsx workbooks are supported", worksheet=worksheet
                    )
            try:
                workbook = load_workbook(stream, read_only=True, data_only=False, keep_links=False)
            except TypeError as exc:
                # Invalid typed style/metadata attributes can raise TypeError during load.
                raise IngestionError(
                    source,
                    "could not read XLSX workbook; invalid workbook attributes",
                    worksheet=worksheet,
                ) from exc
        except IngestionError:
            raise
        except (
            BadZipFile,
            InvalidFileException,
            ParseError,
            OSError,
            EOFError,
            KeyError,
            ValueError,
            IndexError,
        ) as exc:
            # Translate only package validation and workbook-loading failures.
            raise IngestionError(
                source,
                "could not read XLSX workbook; use an unencrypted, valid .xlsx export",
                worksheet=worksheet,
            ) from exc
        try:
            yield workbook
        finally:
            workbook.close()


def inspect_xlsx_sheets(data: bytes, *, source: Source) -> tuple[str, ...]:
    """List ordinary worksheets, including hidden sheets, in stored order."""
    with _open(data, source) as workbook:
        sheets = tuple(sheet.title for sheet in workbook.worksheets)
        if not sheets:
            raise IngestionError(source, "workbook contains no ordinary worksheets")
        return sheets


def _sheet(workbook: Workbook, worksheet: str, source: Source) -> ReadOnlyWorksheet:
    for sheet in workbook.worksheets:
        if sheet.title == worksheet:
            # Ignore declared used-range dimensions; stream the actual worksheet XML.
            # This also handles exporters that incorrectly declare A1:A1.
            sheet.reset_dimensions()
            return sheet
    raise IngestionError(source, "selected worksheet does not exist", worksheet=worksheet)


def _iter_rows(
    sheet: ReadOnlyWorksheet, source: Source, worksheet: str
) -> Iterator[tuple[Cell, ...]]:
    # Read-only worksheets parse lazily. Catch parser failures only while advancing
    # openpyxl's iterator, so errors in the caller's processing still propagate.
    rows = sheet.iter_rows()
    try:
        while True:
            try:
                cells = next(rows)
            except StopIteration:
                return
            except (
                BadZipFile,
                ParseError,
                OSError,
                EOFError,
                KeyError,
                ValueError,
                IndexError,
            ) as exc:
                raise IngestionError(
                    source,
                    "could not read XLSX workbook; use an unencrypted, valid .xlsx export",
                    worksheet=worksheet,
                ) from exc
            yield cells
    finally:
        rows.close()


def _header(rows: Iterator[tuple[Cell, ...]], source: Source, worksheet: str) -> tuple[str, ...]:
    cells = list(next(rows, ()))
    while cells and cells[-1].value is None:
        cells.pop()
    if not cells:
        raise IngestionError(source, "missing header in row 1", source_row=1, worksheet=worksheet)
    names = []
    for index, cell in enumerate(cells, start=1):
        value = cell.value
        if cell.data_type in ("f", "e") or not isinstance(value, str) or not value.strip():
            raise IngestionError(
                source,
                "header must contain nonblank text column names",
                source_row=1,
                column=f"column {index}",
                value=value,
                worksheet=worksheet,
            )
        if value in names:
            raise IngestionError(
                source,
                "duplicate column name",
                source_row=1,
                column=value,
                value=value,
                worksheet=worksheet,
            )
        names.append(value)
    return tuple(names)


def inspect_xlsx_columns(data: bytes, *, source: Source, worksheet: str) -> tuple[str, ...]:
    """Validate row 1 only, with exact text headers and no display-format inference."""
    with _open(data, source, worksheet) as workbook:
        rows = _iter_rows(_sheet(workbook, worksheet, source), source, worksheet)
        try:
            return _header(rows, source, worksheet)
        finally:
            rows.close()


def _numeric(value: int | float) -> Decimal:
    # Never construct Decimal from a float itself or from Excel's number format.
    if isinstance(value, float) and not isfinite(value):
        raise ValueError("numeric value must be finite")
    return Decimal(str(value))


def _evidence(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, str):
        return value
    if isinstance(value, int | float):
        return str(_numeric(value))
    if isinstance(value, datetime | date | time):
        return value.isoformat()
    if isinstance(value, timedelta):
        # Duration-formatted cells are evidence only; no locale-dependent display.
        return str(value)
    raise ValueError(f"unsupported cell value type: {type(value).__name__}")


def _key(value: object) -> str:
    if isinstance(value, str):
        if not value.strip():
            raise ValueError("required key component is blank after trimming")
        return value.strip()
    if isinstance(value, bool) or value is None:
        raise ValueError("key must be nonblank text, a finite number, or a date/time")
    if isinstance(value, int | float):
        number = _numeric(value)
        text = format(number, "f")
        if "." in text:
            text = text.rstrip("0").rstrip(".")
        return "0" if number.is_zero() else text
    if isinstance(value, datetime | date | time):
        return value.isoformat()
    raise ValueError("key must be nonblank text, a finite number, or a date/time")


def _amount(value: object) -> Decimal:
    if isinstance(value, bool) or value is None:
        raise ValueError(
            "amount must be monetary text or a finite number; blank/boolean is invalid"
        )
    if isinstance(value, int | float):
        return _numeric(value)
    if isinstance(value, str):
        return parse_amount(value)
    raise ValueError("amount must be monetary text or a finite number")


def ingest_xlsx(
    data: bytes,
    *,
    source: Source,
    worksheet: str,
    key_columns: Sequence[str],
    amount_column: str,
) -> tuple[SourceRecord, ...]:
    """Validate one sheet atomically; source_row is the actual worksheet row.

    Trailing wholly empty rows are ignored. A wholly empty interior row is an
    error. Selected formulas/errors/blanks are rejected. Evidence is cell-value
    text, not Excel's displayed formatting; unselected formulas remain text.
    """
    if not isinstance(key_columns, Sequence) or isinstance(key_columns, str):
        raise IngestionError(
            source, "key_columns must be a sequence of column names", worksheet=worksheet
        )
    keys = tuple(key_columns)
    if not keys or any(not isinstance(name, str) or not name.strip() for name in keys):
        raise IngestionError(
            source, "select one or more nonblank key column names", worksheet=worksheet
        )
    if len(set(keys)) != len(keys):
        raise IngestionError(source, "selected key columns must be distinct", worksheet=worksheet)
    if not isinstance(amount_column, str) or not amount_column.strip():
        raise IngestionError(source, "select a nonblank amount column name", worksheet=worksheet)
    records = []
    with _open(data, source, worksheet) as workbook:
        rows = _iter_rows(_sheet(workbook, worksheet, source), source, worksheet)
        try:
            header = _header(rows, source, worksheet)
            for column in (*keys, amount_column):
                if column not in header:
                    raise IngestionError(
                        source,
                        "selected column is missing",
                        source_row=1,
                        column=column,
                        worksheet=worksheet,
                    )
            first_blank = None
            for source_row, cells in enumerate(rows, start=2):
                if all(cell.value is None for cell in cells):
                    if first_blank is None:
                        first_blank = source_row
                    continue
                if first_blank is not None:
                    raise IngestionError(
                        source,
                        "wholly blank interior data row",
                        source_row=first_blank,
                        worksheet=worksheet,
                    )
                if any(cell.value is not None for cell in cells[len(header) :]):
                    raise IngestionError(
                        source,
                        "data extends beyond the header columns",
                        source_row=source_row,
                        worksheet=worksheet,
                    )
                fields = dict(zip(header, cells, strict=False))
                converted_keys = []
                for index, column in enumerate((*keys, amount_column)):
                    cell = fields.get(column)
                    value = cell.value if cell is not None else None
                    try:
                        if cell is not None and cell.data_type == "f":
                            raise ValueError(
                                "formula in selected field; export/paste as values first"
                            )
                        if cell is not None and cell.data_type == "e":
                            raise ValueError("Excel error in selected field")
                        if index < len(keys):
                            converted_keys.append(_key(value))
                        else:
                            amount = _amount(value)
                    except ValueError as exc:
                        raise IngestionError(
                            source,
                            str(exc),
                            source_row=source_row,
                            column=column,
                            value=value,
                            worksheet=worksheet,
                        ) from None
                raw_fields = {}
                for column in header:
                    cell = fields.get(column)
                    value = cell.value if cell is not None else None
                    try:
                        raw_fields[column] = _evidence(value)
                    except ValueError as exc:
                        raise IngestionError(
                            source,
                            str(exc),
                            source_row=source_row,
                            column=column,
                            value=value,
                            worksheet=worksheet,
                        ) from None
                records.append(
                    SourceRecord(
                        source,
                        source_row,
                        tuple(converted_keys),
                        amount,
                        raw_fields,
                    )
                )
        finally:
            rows.close()
    return tuple(records)
