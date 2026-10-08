from collections import Counter
from datetime import date, datetime, time
from decimal import Decimal, localcontext
from io import BytesIO, StringIO
from pathlib import Path
from zipfile import ZipFile

import pytest
from openpyxl.styles import Font

from tallydiff import (
    ColumnMapping,
    FindingCategory,
    IngestionError,
    Source,
    export_exceptions_csv,
    export_mapping_profile,
    ingest_csv,
    ingest_xlsx,
    inspect_xlsx_columns,
    inspect_xlsx_sheets,
    load_mapping_profile,
    reconcile,
)


def ingest(data, **kwargs):
    return ingest_xlsx(
        data,
        source=Source.B,
        worksheet="Data",
        key_columns=["id"],
        amount_column="amount",
        **kwargs,
    )


def rewrite_xml(data, path, transform):
    with BytesIO() as output, ZipFile(BytesIO(data)) as original:
        with ZipFile(output, "w") as rewritten:
            for name in original.namelist():
                content = original.read(name)
                if name == path:
                    mutated = transform(content)
                    assert mutated != content, "fixture XML mutation must change the input"
                    content = mutated
                rewritten.writestr(name, content)
        return output.getvalue()


def test_sheet_names_order_hidden_and_single_sheet(xlsx_bytes):
    one = xlsx_bytes({" Export é ": [["id", "amount"]]})
    assert inspect_xlsx_sheets(one, source=Source.A) == (" Export é ",)
    many = xlsx_bytes(
        {"Second": [["id", "amount"]], " First é ": [["ref", "gross"]]},
        configure=lambda book: setattr(book[" First é "], "sheet_state", "hidden"),
    )
    assert inspect_xlsx_sheets(many, source=Source.B) == ("Second", " First é ")
    assert inspect_xlsx_columns(many, source=Source.B, worksheet=" First é ") == ("ref", "gross")
    with pytest.raises(IngestionError, match="does not exist"):
        inspect_xlsx_columns(many, source=Source.A, worksheet="First é")


@pytest.mark.parametrize(
    "header,reason",
    [
        (["id", "amount", "id"], "duplicate"),
        (["id", None, "amount"], "nonblank text"),
        (["id", "  ", "amount"], "nonblank text"),
        ([123, "amount"], "nonblank text"),
        ([date(2026, 1, 2), "amount"], "nonblank text"),
        (["=1+1", "amount"], "nonblank text"),
        ([], "missing header"),
    ],
)
def test_header_validation(xlsx_bytes, header, reason):
    with pytest.raises(IngestionError, match=reason) as error:
        inspect_xlsx_columns(xlsx_bytes({"Data": [header]}), source=Source.A, worksheet="Data")
    assert error.value.source_row == 1
    assert error.value.worksheet == "Data"


def test_exact_header_names_header_only_and_missing_selected_column(xlsx_bytes):
    data = xlsx_bytes({"Data": [[" id ", "amount"]]})
    assert inspect_xlsx_columns(data, source=Source.A, worksheet="Data") == (" id ", "amount")
    assert (
        ingest_xlsx(
            data, source=Source.A, worksheet="Data", key_columns=[" id "], amount_column="amount"
        )
        == ()
    )
    with pytest.raises(IngestionError, match="selected column is missing"):
        ingest(data)


@pytest.mark.parametrize(
    "value,expected",
    [
        ("  000123  ", "000123"),
        ("a  B-1", "a  B-1"),
        (123, "123"),
        (123.0, "123"),
        (123.25, "123.25"),
        (1e-7, "0.0000001"),
        (date(2026, 1, 2), "2026-01-02"),
        (datetime(2026, 1, 2, 3, 4, 5), "2026-01-02T03:04:05"),
        (time(3, 4, 5), "03:04:05"),
    ],
)
def test_key_conversion(xlsx_bytes, value, expected):
    record = ingest(xlsx_bytes({"Data": [["id", "amount"], [value, 1]]}))[0]
    assert record.key == (expected,)
    if isinstance(value, str):
        assert record.raw_fields["id"] == value


def test_numeric_key_format_does_not_reconstruct_leading_zeroes(xlsx_bytes):
    data = xlsx_bytes(
        {"Data": [["id", "amount"], [123, 1], ["000123", 2]]},
        configure=lambda book: setattr(book["Data"]["A2"], "number_format", "000000"),
    )
    assert [row.key for row in ingest(data)] == [("123",), ("000123",)]


@pytest.mark.parametrize(
    "value,expected",
    [
        (" $1,200.00 ", "1200.00"),
        ("(42.50)", "-42.50"),
        (100, "100"),
        (100.01, "100.01"),
        (-42.5, "-42.5"),
        (0, "0"),
        (0.1, "0.1"),
        (1234.12345678901, "1234.12345678901"),
        (1e-12, "0.000000000001"),
    ],
)
def test_amounts_use_decimal_text_never_binary_float_construction(xlsx_bytes, value, expected):
    data = xlsx_bytes(
        {"Data": [["id", "amount"], ["INV", value]]},
        configure=lambda book: setattr(book["Data"]["B2"], "number_format", "$#,##0.00"),
    )
    with localcontext() as context:
        context.prec = 2
        record = ingest(data)[0]
    assert record.amount == Decimal(expected)
    if value == 100.01:
        assert str(record.amount) == "100.01"
    if value == 100:
        assert str(record.amount) == "100"


@pytest.mark.parametrize(
    "column,value,reason",
    [
        ("id", None, "nonblank text"),
        ("id", "  ", "blank after trimming"),
        ("amount", None, "blank/boolean"),
        ("amount", "", "blank/boolean"),
        ("id", True, "nonblank text"),
        ("amount", False, "blank/boolean"),
        ("id", "=1+1", "formula"),
        ("amount", "=100+1", "formula"),
        ("id", "#N/A", "Excel error"),
        ("amount", "#DIV/0!", "Excel error"),
        ("amount", date(2026, 1, 2), "monetary text"),
        ("amount", "NaN", "amount"),
    ],
)
def test_selected_invalid_fields_have_full_context(xlsx_bytes, column, value, reason):
    row = [value, 1] if column == "id" else ["INV", value]
    data = xlsx_bytes({"Data": [["id", "amount"], row]})
    with pytest.raises(IngestionError, match=reason) as error:
        ingest(data)
    assert error.value.source is Source.B
    assert error.value.worksheet == "Data"
    assert error.value.source_row == 2
    assert error.value.column == column
    assert "File B, worksheet 'Data', worksheet row 2" in str(error.value)


@pytest.mark.parametrize("column", ["A", "B"])
def test_nonfinite_numeric_cell_rejected(xlsx_bytes, column):
    data = xlsx_bytes({"Data": [["id", "amount"], [123, 100.01]]})
    data = rewrite_xml(
        data,
        "xl/worksheets/sheet1.xml",
        lambda xml: xml.replace(
            (b"<v>123</v>" if column == "A" else b"<v>100.01</v>"), b"<v>1e309</v>"
        ),
    )
    with pytest.raises(IngestionError, match="finite") as error:
        ingest(data)
    assert error.value.source_row == 2


def test_cell_value_evidence_is_deterministic_and_immutable(xlsx_bytes):
    data = xlsx_bytes(
        {
            "Data": [
                ["id", "amount", "formula", "date", "datetime", "bool", "blank", "error", "memo"],
                [
                    " 000123 ",
                    1234.5,
                    "=SUM(B2:B2)",
                    date(2026, 1, 2),
                    datetime(2026, 1, 2, 3, 4, 5),
                    True,
                    None,
                    "#N/A",
                    'Café, "memo"\nnext',
                ],
            ]
        }
    )
    record = ingest(data)[0]
    assert dict(record.raw_fields) == {
        "id": " 000123 ",
        "amount": "1234.5",
        "formula": "=SUM(B2:B2)",
        "date": "2026-01-02",
        "datetime": "2026-01-02T03:04:05",
        "bool": "TRUE",
        "blank": "",
        "error": "#N/A",
        "memo": 'Café, "memo"\nnext',
    }
    assert ingest(data) == (record,)
    assert record.source_row == 2
    with pytest.raises(TypeError):
        record.raw_fields["id"] = "changed"


def test_trailing_formatted_rows_and_columns_ignored(xlsx_bytes):
    def format_trailing(book):
        book["Data"]["B5000"].number_format = "0.00"
        book["Data"]["Z1"].font = Font(bold=True)

    data = xlsx_bytes({"Data": [["id", "amount"], ["INV", 1]]}, configure=format_trailing)
    assert inspect_xlsx_columns(data, source=Source.A, worksheet="Data") == ("id", "amount")
    assert [row.source_row for row in ingest(data)] == [2]


def test_interior_blank_row_rejected_but_no_partial_records_returned(xlsx_bytes):
    data = xlsx_bytes({"Data": [["id", "amount"], ["A", 1], [], ["B", 2]]})
    with pytest.raises(IngestionError, match="blank interior") as error:
        ingest(data)
    assert error.value.source_row == 3


def test_partial_row_and_extra_column_are_not_silently_dropped(xlsx_bytes):
    for row, reason in [(["A"], "blank/boolean"), (["A", 1, "extra"], "beyond the header")]:
        with pytest.raises(IngestionError, match=reason):
            ingest(xlsx_bytes({"Data": [["id", "amount"], row]}))


def test_incorrect_declared_dimensions_do_not_hide_records(xlsx_bytes):
    data = xlsx_bytes({"Data": [["id", "amount"], ["A", 1], ["B", 2]]})
    data = rewrite_xml(
        data,
        "xl/worksheets/sheet1.xml",
        lambda xml: xml.replace(b'ref="A1:B3"', b'ref="A1:A1"'),
    )
    assert [row.source_row for row in ingest(data)] == [2, 3]


@pytest.mark.parametrize("data", [b"not a workbook", b"PK\x03\x04broken"])
def test_corrupt_workbook_is_contextual(data):
    with pytest.raises(IngestionError, match="could not read XLSX workbook") as error:
        ingest(data)
    assert error.value.worksheet == "Data"


def test_corrupt_sheet_xml_and_missing_zip_part_are_contextual(xlsx_bytes):
    data = xlsx_bytes({"Data": [["id", "amount"], ["A", 1]]})
    broken = rewrite_xml(data, "xl/worksheets/sheet1.xml", lambda xml: b"<broken")
    with pytest.raises(IngestionError, match="could not read XLSX workbook"):
        ingest(broken)
    with BytesIO() as stream:
        with ZipFile(stream, "w") as archive:
            archive.writestr("dummy", "empty")
        with pytest.raises(IngestionError, match="could not read XLSX workbook"):
            inspect_xlsx_sheets(stream.getvalue(), source=Source.A)


def test_resources_closed_on_success_and_validation_failure(xlsx_bytes, monkeypatch):
    import tallydiff.xlsx as module

    opened = []
    original = module.load_workbook

    def capture(*args, **kwargs):
        assert kwargs == {"read_only": True, "data_only": False, "keep_links": False}
        workbook = original(*args, **kwargs)
        opened.append(workbook)
        return workbook

    monkeypatch.setattr(module, "load_workbook", capture)
    ingest(xlsx_bytes({"Data": [["id", "amount"], ["A", 1]]}))
    with pytest.raises(IngestionError):
        ingest(xlsx_bytes({"Data": [["id", "amount"], ["A", "=1"]]}))
    assert all(workbook._archive.fp is None for workbook in opened)


@pytest.mark.parametrize("formats", [("xlsx", "xlsx"), ("csv", "xlsx"), ("xlsx", "csv")])
def test_245_acceptance_all_excel_combinations(acceptance_workbooks, formats):
    mapping = ColumnMapping(
        (("Vendor ID", "Supplier"), ("Invoice Number", "Invoice Ref")),
        "Invoice Amount",
        "Gross Amount",
    )
    profile = load_mapping_profile(export_mapping_profile(mapping))
    inputs = []
    for index, (kind, source, keys, amount, worksheet) in enumerate(
        [
            (formats[0], Source.A, mapping.keys_a, mapping.amount_a, "Ledger A"),
            (formats[1], Source.B, mapping.keys_b, mapping.amount_b, "Ledger B"),
        ]
    ):
        if kind == "xlsx":
            data = acceptance_workbooks[index]
            columns = inspect_xlsx_columns(data, source=source, worksheet=worksheet)
            assert set((*keys, amount)) <= set(columns)
            inputs.append(
                ingest_xlsx(
                    data, source=source, worksheet=worksheet, key_columns=keys, amount_column=amount
                )
            )
        else:
            text = (
                Path(__file__).resolve().parents[1]
                / "sample_data"
                / f"file_{source.value.lower()}.csv"
            ).read_text(encoding="utf-8-sig")
            inputs.append(ingest_csv(text, source=source, key_columns=keys, amount_column=amount))
    profile.validate_columns(tuple(inputs[0][0].raw_fields), tuple(inputs[1][0].raw_fields))
    result = reconcile(*inputs)
    assert (result.total_a, result.total_b, result.control_difference) == (2550, 2305, 245)
    assert result.finding_delta_sum == 245
    assert Counter(f.category for f in result.findings) == {
        FindingCategory.EXACT_MATCH: 1,
        FindingCategory.AMOUNT_MISMATCH: 1,
        FindingCategory.A_ONLY: 1,
        FindingCategory.B_ONLY: 1,
    }
    import csv

    exported = list(csv.DictReader(StringIO(export_exceptions_csv(result).decode("utf-8"))))
    assert [row["Delta (A - B)"] for row in exported] == ["+45", "+500", "-300"]
    assert [row.source_row for row in inputs[0]] == [2, 3, 4]


def test_tolerance_and_duplicates_use_unchanged_engine(xlsx_bytes):
    a = xlsx_bytes({"Data": [["id", "amount"], ["A", 100], ["DUP", 1], ["DUP", 2]]})
    b = xlsx_bytes({"Data": [["id", "amount"], ["A", 100.01], ["DUP", 3]]})
    records_a = ingest_xlsx(
        a, source=Source.A, worksheet="Data", key_columns=["id"], amount_column="amount"
    )
    result = reconcile(records_a, ingest(b), amount_tolerance=Decimal("0.01"))
    assert [f.category for f in result.findings] == [
        FindingCategory.WITHIN_TOLERANCE,
        FindingCategory.DUPLICATE_AMBIGUOUS,
    ]
    assert result.tolerated_delta_total == result.control_difference == Decimal("-0.01")
    assert len(result.exceptions[0].rows_a) == 2
    assert result.finding_delta_sum == result.control_difference


@pytest.mark.parametrize(
    "content_type",
    [
        b"application/vnd.ms-excel.sheet.macroEnabled.main+xml",
        b"application/vnd.openxmlformats-officedocument.spreadsheetml.template.main+xml",
    ],
)
def test_other_workbook_package_types_are_rejected_even_if_renamed(xlsx_bytes, content_type):
    data = xlsx_bytes({"Data": [["id", "amount"], ["A", 1]]})
    data = rewrite_xml(
        data,
        "[Content_Types].xml",
        lambda xml: xml.replace(
            b"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml",
            content_type,
        ),
    )
    with pytest.raises(IngestionError, match="only standard .xlsx"):
        ingest(data)


def test_unexpected_programming_failure_is_not_swallowed(xlsx_bytes, monkeypatch):
    def fail(*args, **kwargs):
        raise RuntimeError("unexpected parser failure")

    monkeypatch.setattr("tallydiff.xlsx.load_workbook", fail)
    with pytest.raises(RuntimeError, match="unexpected parser failure"):
        ingest(xlsx_bytes({"Data": [["id", "amount"], ["A", 1]]}))


def test_invalid_typed_style_attribute_has_workbook_read_error(xlsx_bytes):
    data = xlsx_bytes({"Data": [["id", "amount"], ["A", 1]]})
    data = rewrite_xml(
        data, "xl/styles.xml", lambda xml: xml.replace(b'sz val="11"', b'sz val="invalid"')
    )
    with pytest.raises(IngestionError, match="could not read XLSX workbook"):
        inspect_xlsx_sheets(data, source=Source.A)


def test_downstream_value_error_propagates_unchanged_and_closes_workbook(xlsx_bytes):
    from tallydiff.xlsx import _open

    data = xlsx_bytes({"Data": [["id", "amount"], ["A", 1]]})
    unexpected = ValueError("unexpected downstream failure")
    with pytest.raises(ValueError) as error:
        with _open(data, Source.B, "Data") as workbook:
            assert workbook["Data"].title == "Data"
            raise unexpected
    assert error.value is unexpected
    assert workbook._archive.fp is None


def test_amount_column_can_also_supply_a_normalized_key(xlsx_bytes):
    data = xlsx_bytes({"Data": [["id", "amount"], [" 000123 ", 1.2]]})
    record = ingest_xlsx(
        data,
        source=Source.A,
        worksheet="Data",
        key_columns=["id", "amount"],
        amount_column="amount",
    )[0]
    assert record.key == ("000123", "1.2")
    assert record.amount == Decimal("1.2")
